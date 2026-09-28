"""openpyxl 直写导出（rows_to_xlsx）与登录限流的单元测试（不连数据库）。"""

import time

import openpyxl
import pytest

from app import auth
from app.utils import rows_to_xlsx

# ---------- rows_to_xlsx ----------

def _load(buf):
    buf.seek(0)
    return openpyxl.load_workbook(buf)


def test_rows_to_xlsx_headers_and_values():
    buf = rows_to_xlsx([('单据', [
        {'单据编号': 'ZC260001', '金额': 1234.5},
        {'单据编号': 'ZC260002', '金额': None},
    ])])
    wb = _load(buf)
    assert wb.sheetnames == ['单据']
    ws = wb['单据']
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0] == ('单据编号', '金额')          # 表头取第一行键序
    assert rows[1] == ('ZC260001', 1234.5)
    assert rows[2] == ('ZC260002', None)            # None → 空白单元格


def test_rows_to_xlsx_multi_sheet_and_empty():
    buf = rows_to_xlsx([('A', [{'x': 1}]), ('B', []), ('C', [])])
    wb = _load(buf)
    assert wb.sheetnames == ['A', 'B', 'C']
    assert list(wb['B'].iter_rows(values_only=True)) == []  # 空数据 → 空表


def test_rows_to_xlsx_column_follows_first_row_keys():
    # 后续行多余键被忽略、缺键写空白（get 语义）
    buf = rows_to_xlsx([('S', [{'a': 1, 'b': 2}, {'a': 3, 'c': 4}])])
    rows = list(_load(buf)['S'].iter_rows(values_only=True))
    assert rows[1] == (1, 2)
    assert rows[2] == (3, None)


# ---------- 登录限流 ----------

@pytest.fixture(autouse=True)
def _clean_throttle_state():
    """隔离模块级限流状态，避免用例间/真实服务状态串扰。"""
    auth._login_failures.clear()
    auth._login_locks.clear()
    yield
    auth._login_failures.clear()
    auth._login_locks.clear()


def test_locks_after_max_failures():
    for _ in range(auth.LOGIN_MAX_FAILURES):
        auth._register_failure('ip|user')
    assert auth._locked_seconds('ip|user') > 0
    assert auth._locked_seconds('ip|user') <= auth.LOGIN_LOCK_SECONDS
    # 计数在锁定后清零：解锁后重新累计（防永久封禁残留）
    assert 'ip|user' not in auth._login_failures


def test_below_threshold_not_locked():
    for _ in range(auth.LOGIN_MAX_FAILURES - 1):
        auth._register_failure('ip|user')
    assert auth._locked_seconds('ip|user') == 0


def test_window_expiry_resets_counter():
    # 窗口外的一次失败 + 窗口内的 4 次 → 不锁定（计数已重置）
    auth._login_failures['ip|user'] = [1, time.monotonic() - auth.LOGIN_WINDOW_SECONDS - 1]
    for _ in range(auth.LOGIN_MAX_FAILURES - 1):
        auth._register_failure('ip|user')
    assert auth._locked_seconds('ip|user') == 0


def test_lock_expiry_clears_state():
    auth._register_failure('ip|user')
    auth._login_locks['ip|user'] = time.monotonic() - 1  # 已过期
    assert auth._locked_seconds('ip|user') == 0
    assert 'ip|user' not in auth._login_locks
    assert 'ip|user' not in auth._login_failures


def test_clear_failures():
    auth._register_failure('ip|user')
    auth._clear_failures('ip|user')
    assert auth._locked_seconds('ip|user') == 0
    assert 'ip|user' not in auth._login_failures


# ---------- 上传白名单 ----------

def test_allowed_extensions():
    from app.inventory import allowed_file
    assert allowed_file('报告.PDF')
    assert allowed_file('a/b/c照片.jpg')
    assert not allowed_file('恶意.exe')
    assert not allowed_file('无后缀')
    assert not allowed_file('')

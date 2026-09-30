# -*- coding: utf-8 -*-
"""对比：用户修改后的 Excel（明细） vs 数据库 inventory 当前记录（DZ2601~04）。

按资产编码对齐，比较部门/使用人/类型/日期/状态/标签；规格仅 Excel 有，
单独汇总。结果写 UTF-8 报告文件。
"""
from datetime import date, datetime

from openpyxl import load_workbook

from app.config import get_db_connection

XLSX = 'scripts/DZ2601-04_asset_changes_export.xlsx'
OUT = 'scripts/dz2601_04_diff_report.txt'
MONTH_PREFIXES = ('DZ2601', 'DZ2602', 'DZ2603', 'DZ2604')

# Excel 列名 -> DB 字段
FIELD_MAP = [
    ('部门', 'department'),
    ('使用人', 'site'),
    ('资产类型', 'type'),
    ('发放日期', 'datetime'),
    ('资产状态', 'status'),
    ('流转标签', 'tag'),
]


def norm(v):
    """单元格/DB 值归一化为可比较字符串。"""
    if v is None:
        return ''
    if isinstance(v, datetime):
        return v.strftime('%Y-%m-%d')
    if isinstance(v, date):
        return v.isoformat()
    return str(v).strip()


def load_excel():
    ws = load_workbook(XLSX)['明细']
    rows = list(ws.iter_rows(values_only=True))
    headers = [str(h).strip() if h is not None else '' for h in rows[0]]
    data = []
    for r in rows[1:]:
        if all(c is None or str(c).strip() == '' for c in r):
            continue
        item = {h: c for h, c in zip(headers, r)}
        num = norm(item.get('资产编码'))
        if num:
            data.append((num, item))
    return data


def load_db():
    conn = get_db_connection()
    if not conn:
        raise SystemExit('DB connect failed')
    cursor = conn.cursor(dictionary=True)
    like_sql = ' OR '.join(['number LIKE %s'] * len(MONTH_PREFIXES))
    params = [f'{p}%' for p in MONTH_PREFIXES]
    cursor.execute(
        f"SELECT number, department, site, type, "
        f"DATE_FORMAT(datetime, '%Y-%m-%d') AS datetime, status, tag "
        f"FROM inventory WHERE {like_sql} ORDER BY number, id",
        params,
    )
    rows = cursor.fetchall()
    cursor.close()
    conn.close()
    # 每个编码理论上一条当前记录；若出现多条取最新
    by_num = {}
    for r in rows:
        by_num[r['number']] = r
    return by_num


def main():
    excel_rows = load_excel()
    db_rows = load_db()

    excel_nums = [n for n, _ in excel_rows]
    dup = {n for n in excel_nums if excel_nums.count(n) > 1}
    excel_by_num = dict(excel_rows)
    db_nums = set(db_rows)

    common = [n for n in excel_nums if n in db_nums]
    only_excel = [n for n in excel_nums if n not in db_nums]
    only_db = sorted(db_nums - set(excel_nums))

    lines = []
    lines.append('=' * 72)
    lines.append('Excel vs 资产变更(inventory) 差异报告')
    lines.append(f'Excel 明细 {len(excel_rows)} 行（{len(set(excel_nums))} 个编码） | '
                 f'数据库 {len(db_rows)} 条记录')
    lines.append('=' * 72)

    if dup:
        lines.append(f'\n【警告】Excel 中重复的资产编码: {", ".join(sorted(dup))}')

    # ---- 字段差异 ----
    diff_cnt = 0
    space_only_cnt = 0
    spec_filled = []
    raw_vs_norm = []
    for num in common:
        item = excel_by_num[num]
        db = db_rows[num]
        diffs = []
        space_only = []
        for xl_col, db_col in FIELD_MAP:
            xv, dv = norm(item.get(xl_col)), norm(db.get(db_col))
            raw_xv = '' if item.get(xl_col) is None else str(item.get(xl_col))
            if raw_xv != raw_xv.strip() and raw_xv.strip() == dv:
                space_only.append(xl_col)
            if xv != dv:
                diffs.append((xl_col, xv, dv))
        if item.get('资产规格') not in (None, ''):
            spec_filled.append(num)
        if space_only:
            space_only_cnt += 1
            raw_vs_norm.append((num, space_only))
        if diffs:
            diff_cnt += 1
            lines.append(f'\n◆ {num}')
            for col, xv, dv in diffs:
                lines.append(f'    {col}: Excel={xv!r}  数据库={dv!r}')

    lines.append('')
    lines.append('-' * 72)
    lines.append(f'字段有实质差异的资产: {diff_cnt} 个（相同编码 {len(common)} 个中）')
    if space_only_cnt:
        lines.append(f'仅前后空格差异的资产: {space_only_cnt} 个 '
                     f'({", ".join(f"{n}[{",".join(c)}]" for n, c in raw_vs_norm)})')
    lines.append(f'Excel 填写了规格的资产: {len(spec_filled)} 个（数据库规格字段均为空，不计入字段差异）')

    if only_excel:
        lines.append(f'\n【仅 Excel 有】（新增或改了编码）{len(only_excel)} 个: '
                     f'{", ".join(only_excel)}')
    if only_db:
        lines.append(f'\n【仅数据库有】（Excel 删除或漏掉）{len(only_db)} 个: '
                     f'{", ".join(only_db)}')

    with open(OUT, 'w', encoding='utf-8') as f:
        f.write('\n'.join(lines) + '\n')
    print('written', OUT)


if __name__ == '__main__':
    main()
